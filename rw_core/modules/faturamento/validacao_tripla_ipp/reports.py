from __future__ import annotations

from copy import copy
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from src.config.settings import DATA_DIR
from rw_core.modules.faturamento.coupa.imports import normalizar_cnpj_cpf
from rw_core.normalizers.fields import parse_number


TEMPLATE_DIR = DATA_DIR / "templates"
REPORT_TEMPLATES = {
    "Coleta": TEMPLATE_DIR / "tripla_ipp_coleta_modelo.xlsx",
    "Transferencia": TEMPLATE_DIR / "tripla_ipp_transferencia_modelo.xlsx",
}
REPORT_SHEET_NAME = "Complemento - Frete"
REPORT_START_ROW = 6
INTEGER_NUMBER_FORMAT = '#,##0'
MONEY_NUMBER_FORMAT = 'R$ #,##0.00'


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    text = str(value).strip()
    return "" if text.upper() in {"NONE", "NAN", "<NA>", "NAT"} else text


def _number(value: Any) -> float | None:
    parsed = parse_number(value)
    if parsed is None:
        return None
    return float(parsed)


def _first_number(row: dict[str, Any], columns: list[str]) -> float | None:
    for column in columns:
        parsed = _number(row.get(column))
        if parsed is not None:
            return parsed
    return None


def _first_text(row: dict[str, Any], columns: list[str]) -> str:
    for column in columns:
        text = _clean_text(row.get(column))
        if text:
            return text
    return ""


def _copy_row_style(ws, source_row: int, target_row: int) -> None:
    if source_row == target_row:
        return
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for column in range(1, ws.max_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)


def _filial_code(record: dict[str, Any], filial_lookup: dict[str, str] | None) -> str:
    if not filial_lookup:
        return ""
    cnpj_cobranca = normalizar_cnpj_cpf(_first_text(record, ["CNPJ/CPF da cobranca", "CNPJ Cobranca", "CNPJ Cobrança"]))
    return filial_lookup.get(cnpj_cobranca, "") if cnpj_cobranca else ""


def _fill_report_row(ws, row_number: int, panel: str, record: dict[str, Any], filial_lookup: dict[str, str] | None) -> None:
    is_coleta = panel == "Coleta"
    nf = _first_text(record, ["NF", "Notas Fiscais"])
    cnpj_origem = _first_text(record, ["CNPJ Origem", "_cnpj_remetente"])
    cnpj_destino = _first_text(record, ["CNPJ Destino", "_cnpj_destinatario"])
    volume = _first_number(record, ["Quantidade Portal", "Volume"])
    km = _first_number(record, ["KM referente origem e destino"])
    peso_frete = _first_number(record, ["Valor LCTE", "Peso frete", "Valor total LCTE"])
    valor_portal_transportador = _first_number(record, ["Valor do frete_1", "Valor Portran", "Valor Calculado Portran"])
    valor_total_portal = valor_portal_transportador if valor_portal_transportador is not None else peso_frete
    coupa_tariff = _first_number(record, ["Tarifa Coupa", "Bitrem (40 a 49m3)", "Rodotrem (maior 50m3)"])
    operation_label = "Coleta" if is_coleta else "Transferência"

    ws.cell(row_number, 1).value = _filial_code(record, filial_lookup)
    ws.cell(row_number, 2).value = operation_label
    ws.cell(row_number, 3).value = 2219
    ws.cell(row_number, 4).value = nf
    ws.cell(row_number, 5).value = ""
    ws.cell(row_number, 6).value = cnpj_origem
    ws.cell(row_number, 7).value = cnpj_destino

    if is_coleta:
        ws.cell(row_number, 8).value = "R$/km"
        ws.cell(row_number, 9).value = "N/A"
        ws.cell(row_number, 10).value = km if km is not None else "N/A"
        ws.cell(row_number, 11).value = f"=L{row_number}/J{row_number}"
        ws.cell(row_number, 12).value = valor_total_portal
        ws.cell(row_number, 13).value = "R$/km"
        ws.cell(row_number, 14).value = "N/A"
        ws.cell(row_number, 15).value = km if km is not None else "N/A"
        ws.cell(row_number, 16).value = coupa_tariff
        ws.cell(row_number, 17).value = f"=P{row_number}*O{row_number}"
    else:
        ws.cell(row_number, 8).value = "R$/Litro"
        ws.cell(row_number, 9).value = volume
        ws.cell(row_number, 10).value = "N/A"
        ws.cell(row_number, 11).value = f"=L{row_number}/I{row_number}"
        ws.cell(row_number, 12).value = valor_total_portal
        ws.cell(row_number, 13).value = "R$/Litro"
        ws.cell(row_number, 14).value = volume
        ws.cell(row_number, 15).value = "N/A"
        ws.cell(row_number, 16).value = coupa_tariff
        ws.cell(row_number, 17).value = f"=P{row_number}*N{row_number}"

    ws.cell(row_number, 18).value = f"=Q{row_number}-L{row_number}"
    for column in [9, 10, 14, 15]:
        ws.cell(row_number, column).number_format = INTEGER_NUMBER_FORMAT
    for column in [11, 12, 16, 17, 18]:
        ws.cell(row_number, column).number_format = MONEY_NUMBER_FORMAT


def generate_standard_report(panel: str, rows: pd.DataFrame, filial_lookup: dict[str, str] | None = None) -> bytes:
    if panel not in REPORT_TEMPLATES:
        raise ValueError(f"Painel sem modelo de relatorio: {panel}")
    template_path = REPORT_TEMPLATES[panel]
    if not Path(template_path).exists():
        raise FileNotFoundError(f"Modelo nao encontrado: {template_path}")
    records = rows.to_dict(orient="records") if isinstance(rows, pd.DataFrame) else []
    if not records:
        raise ValueError("Selecione ao menos uma linha para gerar o relatorio.")

    wb = load_workbook(template_path)
    ws = wb[REPORT_SHEET_NAME] if REPORT_SHEET_NAME in wb.sheetnames else wb.active
    for offset, record in enumerate(records):
        row_number = REPORT_START_ROW + offset
        if offset:
            ws.insert_rows(row_number)
            _copy_row_style(ws, REPORT_START_ROW, row_number)
        _fill_report_row(ws, row_number, panel, record, filial_lookup)

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass

    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()
